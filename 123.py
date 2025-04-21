import re
import openpyxl

# File names
input_file = 'input.txt'
output_file = 'output.xlsx'

# Create Excel workbook and sheet
workbook = openpyxl.Workbook()
sheet = workbook.active
sheet.title = '12 Digit Numbers'

# Add headers
sheet['A1'] = '12-Digit Number'
sheet['B1'] = 'Line Found'

# Regex: 12 digits, not part of a larger number
pattern = r'(?<!\d)(\d{12})(?!\d)'

# Row counter
row = 2

# Read file and find matches
with open(input_file, 'r', encoding='utf-8') as file:
    for line in file:
        matches = re.findall(pattern, line)
        for match in matches:
            sheet.cell(row=row, column=1).value = match
            sheet.cell(row=row, column=2).value = line.strip()
            row += 1

# Save results
workbook.save(output_file)
print(f"Done! 12-digit numbers saved in '{output_file}'.")
